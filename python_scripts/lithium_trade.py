#!/usr/bin/env python3
"""
Pull HS 6-digit lithium trade data (imports + exports) for all available years.

DATA SOURCE
-----------
World Integrated Trade Solution (WITS) serves its HS 6-digit product-level trade
data from UN Comtrade -- its own download pages are labelled "WITS - UNSD
Comtrade". WITS's public API only exposes aggregated product *groups*
(e.g. "84-85_MachElec") and carries no quantity or unit, so this script queries
the same underlying series directly from the UN Comtrade API. The numbers match
what WITS shows under Trade Stats > By Product (HS 6-digit).

OUTPUT
------
  output/lithium_imports.csv
  output/lithium_exports.csv
  output/lithium_trade.xlsx   (two sheets: Imports, Exports -- needs openpyxl)

CONFIGURE
---------
Edit the CONFIG block below. See README_lithium_trade.md for details.

USAGE
-----
  python3 lithium_trade.py
  python3 lithium_trade.py --years 2015-2024 --partner all
  python3 lithium_trade.py --list-codes lithium
"""

from __future__ import annotations

import argparse
import csv
import datetime as dt
import hashlib
import json
import os
import pathlib
import sys
import time

try:
    import requests
except ImportError:
    sys.exit("Missing dependency. Run:  pip install requests")


# ===========================================================================
# CONFIG -- edit this block
# ===========================================================================

# HS 6-digit product codes to pull, mapped to your own label for each.
# The label is only cosmetic; the code is what gets queried. Comment lines
# out to drop them, add lines to include more. Use --list-codes to search.
PRODUCT_CODES = {
    "253090": "Lithium ores / spodumene (mineral substances n.e.c., ch. 25)",
    "280519": "Lithium metal (alkali metals other than sodium & calcium)",
    "282520": "Lithium oxide and hydroxide",
    "283691": "Lithium carbonate",
    "850650": "Primary cells and batteries, lithium",
    "850760": "Lithium-ion accumulators (batteries)",
}

# Years to pull.
#   "all"            -> every year the classification exists (1988 -> present)
#   (2010, 2024)     -> inclusive range
#   [2019, 2021]     -> explicit list
YEARS = "all"

# Reporting countries.  "all" -> every reporter, or a list of ISO3 codes
# e.g. ["USA", "CHN", "AUS", "CHL"]
REPORTERS = "all"

# Trading partner.
#   "world" -> one aggregate row per reporter (recommended; compact)
#   "all"   -> bilateral detail, one row per partner (MUCH bigger + slower)
#   ["CHN", "USA"] -> specific partners by ISO3
PARTNER = "world"

# Trade flows to pull, mapped to output sheet/file names.
FLOWS = {"M": "Imports", "X": "Exports"}

# Classification. "HS" follows each country's own reported HS revision, which
# is what WITS does. Use "H0".."H6" to force a single revision.
CLASSIFICATION = "HS"

OUTPUT_DIR = "output"

# Seconds to wait between API calls. The keyless endpoint rate-limits
# aggressively (HTTP 429); 6s is a safe default. With an API key set
# (see below), this drops automatically.
REQUEST_DELAY_SECONDS = 6.0

# Cache raw API responses so reruns and interrupted runs are cheap.
# Delete the cache directory to force a fresh pull.
USE_CACHE = True
CACHE_DIR = ".comtrade_cache"

# Optional. A free UN Comtrade subscription key removes the 500-row cap
# (raising it to 100,000) and the tight rate limit, cutting a full run from
# roughly an hour to a couple of minutes. Register at
# https://comtradedeveloper.un.org/ then:  export COMTRADE_API_KEY=...
API_KEY_ENV_VAR = "COMTRADE_API_KEY"

# ===========================================================================
# End of CONFIG
# ===========================================================================


BASE_PUBLIC = "https://comtradeapi.un.org/public/v1/preview/C/A"
BASE_KEYED = "https://comtradeapi.un.org/data/v1/get/C/A"
REFERENCE = "https://comtradeapi.un.org/files/v1/app/reference"

PREVIEW_ROW_CAP = 500      # keyless endpoint truncates silently at 500 rows
KEYED_ROW_CAP = 100_000
FIRST_HS_YEAR = 1988       # HS reporting begins in 1988

OUTPUT_COLUMNS = [
    "country",
    "country_iso3",
    "country_code",
    "year",
    "trade_flow",
    "product_code",
    "product_label",
    "product_description",
    "trade_value_usd",
    "quantity",
    "unit",
    "unit_code",
    "net_weight_kg",
    "partner",
    "partner_iso3",
    "quantity_is_estimated",
    "hs_revision",
    "value_basis",
]


# --------------------------------------------------------------------------
# HTTP
# --------------------------------------------------------------------------

class Client:
    """UN Comtrade client with caching, throttling and 429 backoff."""

    def __init__(self, api_key: str | None, delay: float, cache_dir: pathlib.Path | None):
        self.api_key = api_key
        self.delay = 0.6 if api_key else delay
        self.row_cap = KEYED_ROW_CAP if api_key else PREVIEW_ROW_CAP
        self.cache_dir = cache_dir
        self.session = requests.Session()
        self.session.headers["User-Agent"] = "lithium-trade-puller/1.0"
        if api_key:
            self.session.headers["Ocp-Apim-Subscription-Key"] = api_key
        self._last_call = 0.0
        self.calls = 0
        self.cache_hits = 0
        if cache_dir:
            cache_dir.mkdir(parents=True, exist_ok=True)

    def _throttle(self) -> None:
        wait = self.delay - (time.monotonic() - self._last_call)
        if wait > 0:
            time.sleep(wait)
        self._last_call = time.monotonic()

    def _cache_path(self, params: dict) -> pathlib.Path | None:
        if not self.cache_dir:
            return None
        key = json.dumps(params, sort_keys=True)
        return self.cache_dir / (hashlib.sha1(key.encode()).hexdigest() + ".json")

    def get_data(self, params: dict) -> list[dict]:
        """One data call. Returns the row list; raises on unrecoverable error."""
        cached = self._cache_path(params)
        if cached and cached.exists():
            self.cache_hits += 1
            return json.loads(cached.read_text())["data"]

        base = BASE_KEYED if self.api_key else BASE_PUBLIC
        url = f"{base}/{CLASSIFICATION}"

        for attempt in range(6):
            self._throttle()
            self.calls += 1
            try:
                resp = self.session.get(url, params=params, timeout=180)
            except requests.RequestException as exc:
                if attempt == 5:
                    raise RuntimeError(f"network error after 6 tries: {exc}") from exc
                time.sleep(2 ** attempt * 5)
                continue

            if resp.status_code == 429 or resp.status_code >= 500:
                retry_after = resp.headers.get("Retry-After")
                nap = float(retry_after) if retry_after and retry_after.isdigit() else 2 ** attempt * 10
                print(f"    rate limited ({resp.status_code}); waiting {nap:.0f}s", flush=True)
                time.sleep(nap)
                continue

            if resp.status_code == 401:
                raise RuntimeError("401 Unauthorized -- the API key in "
                                   f"${API_KEY_ENV_VAR} was rejected.")
            if resp.status_code != 200:
                raise RuntimeError(f"HTTP {resp.status_code}: {resp.text[:300]}")

            payload = resp.json()
            if payload.get("error"):
                raise RuntimeError(f"API error: {payload['error']}")
            rows = payload.get("data") or []
            if cached:
                cached.write_text(json.dumps({"data": rows}))
            return rows

        raise RuntimeError("gave up after repeated rate limiting -- rerun later "
                           "(cached progress is kept) or set an API key")

    def get_reference(self, name: str) -> list[dict]:
        cached = self._cache_path({"ref": name})
        if cached and cached.exists():
            return json.loads(cached.read_text())["data"]
        resp = self.session.get(f"{REFERENCE}/{name}.json", timeout=120)
        resp.raise_for_status()
        payload = resp.json()
        rows = payload["results"] if isinstance(payload, dict) else payload
        if cached:
            cached.write_text(json.dumps({"data": rows}))
        return rows


# --------------------------------------------------------------------------
# Reference lookups
# --------------------------------------------------------------------------

class Lookups:
    def __init__(self, client: Client):
        print("Loading reference tables...", flush=True)
        self.reporters = {}
        self.reporter_by_iso3 = {}
        for r in client.get_reference("Reporters"):
            code = int(r["reporterCode"])
            iso3 = (r.get("reporterCodeIsoAlpha3") or "").strip()
            self.reporters[code] = {
                "name": r.get("reporterDesc") or r.get("text") or str(code),
                "iso3": iso3,
                "is_group": bool(r.get("isGroup")),
            }
            if iso3:
                self.reporter_by_iso3[iso3.upper()] = code

        self.partners = {}
        self.partner_by_iso3 = {}
        for p in client.get_reference("partnerAreas"):
            try:
                code = int(p["PartnerCode"] if "PartnerCode" in p else p.get("id"))
            except (TypeError, ValueError):
                continue
            iso3 = (p.get("PartnerCodeIsoAlpha3") or p.get("partnerCodeIsoAlpha3") or "").strip()
            self.partners[code] = {
                "name": p.get("PartnerDesc") or p.get("text") or str(code),
                "iso3": iso3,
            }
            if iso3:
                self.partner_by_iso3.setdefault(iso3.upper(), code)

        self.units = {}
        for u in client.get_reference("QuantityUnits"):
            self.units[int(u["qtyCode"])] = {
                "abbr": u.get("qtyAbbr") or "",
                "desc": u.get("qtyDescription") or "",
            }

        # Commodity descriptions: merge every HS revision so old codes resolve.
        self.commodities = {}
        for rev in ("H0", "H1", "H2", "H3", "H4", "H5", "H6"):
            try:
                for c in client.get_reference(rev):
                    cid = str(c.get("id", "")).strip()
                    text = str(c.get("text", ""))
                    if cid and cid not in self.commodities:
                        # entries look like "283691 - Carbonates; lithium carbonate"
                        self.commodities[cid] = text.split(" - ", 1)[-1].strip()
            except Exception:
                continue

    def reporter_codes(self, spec) -> list[int]:
        if spec == "all":
            return sorted(c for c, v in self.reporters.items()
                          if not v["is_group"] and c != 0)
        codes = []
        for item in spec:
            key = str(item).upper()
            if key.isdigit():
                codes.append(int(key))
            elif key in self.reporter_by_iso3:
                codes.append(self.reporter_by_iso3[key])
            else:
                raise SystemExit(f"Unknown reporter: {item!r} (use ISO3, e.g. 'USA')")
        return sorted(set(codes))

    def partner_param(self, spec) -> str:
        if spec == "world":
            return "0"
        if spec == "all":
            return ""
        codes = []
        for item in spec:
            key = str(item).upper()
            if key.isdigit():
                codes.append(int(key))
            elif key in self.partner_by_iso3:
                codes.append(self.partner_by_iso3[key])
            else:
                raise SystemExit(f"Unknown partner: {item!r} (use ISO3, e.g. 'CHN')")
        return ",".join(str(c) for c in sorted(set(codes)))


# --------------------------------------------------------------------------
# Fetching
# --------------------------------------------------------------------------

def fetch_partition(client: Client, year: int, flow: str, codes: list[str],
                    reporters: list[int] | None, partner_param: str,
                    depth: int = 0) -> list[dict]:
    """
    Fetch one slice, splitting recursively when the API row cap is hit.

    The keyless endpoint truncates at 500 rows without saying so, so any
    response at or above the cap is treated as incomplete and subdivided
    until every piece comes back under it.
    """
    params = {
        "reporterCode": ",".join(str(r) for r in reporters) if reporters else "",
        "period": str(year),
        "partnerCode": partner_param,
        "cmdCode": ",".join(codes),
        "flowCode": flow,
    }
    rows = client.get_data(params)

    if len(rows) < client.row_cap:
        return rows

    indent = "    " + "  " * depth
    # Split on whichever dimension still has room to divide.
    if len(codes) > 1:
        mid = len(codes) // 2
        print(f"{indent}row cap hit -- splitting {len(codes)} product codes", flush=True)
        return (fetch_partition(client, year, flow, codes[:mid], reporters, partner_param, depth + 1)
                + fetch_partition(client, year, flow, codes[mid:], reporters, partner_param, depth + 1))

    if reporters is None or len(reporters) > 1:
        if reporters is None:
            raise RuntimeError(
                f"{year} {flow} {codes}: row cap hit with an unbounded reporter list. "
                "Set REPORTERS to an explicit list, or use an API key.")
        mid = len(reporters) // 2
        print(f"{indent}row cap hit -- splitting {len(reporters)} reporters", flush=True)
        return (fetch_partition(client, year, flow, codes, reporters[:mid], partner_param, depth + 1)
                + fetch_partition(client, year, flow, codes, reporters[mid:], partner_param, depth + 1))

    print(f"{indent}WARNING: {year} {flow} {codes} reporter={reporters} still at the "
          f"{client.row_cap}-row cap; results for this slice may be truncated.", flush=True)
    return rows


def normalise(raw: list[dict], lookups: Lookups, flow_label: str) -> list[dict]:
    out = []
    for r in raw:
        rep_code = r.get("reporterCode")
        rep = lookups.reporters.get(rep_code, {})
        par_code = r.get("partnerCode")
        par = lookups.partners.get(par_code, {})
        unit_code = r.get("qtyUnitCode")
        unit = lookups.units.get(unit_code, {}) if unit_code is not None else {}
        code = str(r.get("cmdCode") or "")
        qty = r.get("qty")
        # Comtrade uses -1 / "no quantity" to mean the reporter filed none.
        if unit_code in (-1, None) or qty in (0, None):
            qty_out = "" if qty in (None, 0) and unit_code in (-1, None) else qty
        else:
            qty_out = qty

        out.append({
            "country": rep.get("name") or r.get("reporterDesc") or rep_code,
            "country_iso3": rep.get("iso3") or (r.get("reporterISO") or ""),
            "country_code": rep_code,
            "year": r.get("refYear") or r.get("period"),
            "trade_flow": flow_label,
            "product_code": code,
            "product_label": PRODUCT_CODES.get(code, ""),
            "product_description": lookups.commodities.get(code, ""),
            "trade_value_usd": r.get("primaryValue"),
            "quantity": qty_out,
            "unit": unit.get("abbr") or ("N/A" if unit_code == -1 else ""),
            "unit_code": unit_code,
            "net_weight_kg": r.get("netWgt") or "",
            "partner": par.get("name") or ("World" if par_code == 0 else par_code),
            "partner_iso3": par.get("iso3") or ("WLD" if par_code == 0 else ""),
            "quantity_is_estimated": r.get("isQtyEstimated"),
            "hs_revision": r.get("classificationCode") or "",
            "value_basis": "CIF" if flow_label == "Imports" else "FOB",
        })
    return out


# --------------------------------------------------------------------------
# Output
# --------------------------------------------------------------------------

def sort_key(row: dict):
    return (row["product_code"], str(row["country"]), row["year"] or 0)


def write_csv(path: pathlib.Path, rows: list[dict]) -> None:
    with path.open("w", newline="", encoding="utf-8-sig") as fh:
        writer = csv.DictWriter(fh, fieldnames=OUTPUT_COLUMNS, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def write_xlsx(path: pathlib.Path, sheets: dict[str, list[dict]]) -> bool:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
        from openpyxl.utils import get_column_letter
    except ImportError:
        return False

    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(title=name)
        ws.append(OUTPUT_COLUMNS)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for row in rows:
            ws.append([row.get(col, "") for col in OUTPUT_COLUMNS])
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(OUTPUT_COLUMNS))}{ws.max_row}"
        for idx, col in enumerate(OUTPUT_COLUMNS, start=1):
            width = max(len(col) + 2, 12)
            if col in ("country", "product_description", "product_label", "partner"):
                width = 34
            ws.column_dimensions[get_column_letter(idx)].width = width
    wb.save(path)
    return True


# --------------------------------------------------------------------------
# Main
# --------------------------------------------------------------------------

def resolve_years(spec) -> list[int]:
    this_year = dt.date.today().year
    if spec == "all":
        return list(range(FIRST_HS_YEAR, this_year + 1))
    if isinstance(spec, tuple) and len(spec) == 2:
        return list(range(int(spec[0]), int(spec[1]) + 1))
    return [int(y) for y in spec]


def parse_years_arg(text: str):
    if text == "all":
        return "all"
    if "-" in text:
        lo, hi = text.split("-", 1)
        return (int(lo), int(hi))
    return [int(y) for y in text.split(",")]


def parse_list_arg(text: str):
    if text in ("all", "world"):
        return text
    return [t.strip().upper() for t in text.split(",") if t.strip()]


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--codes", help="Comma-separated HS6 codes, overriding PRODUCT_CODES")
    ap.add_argument("--years", help="'all', '2015-2024', or '2019,2021'")
    ap.add_argument("--reporters", help="'all' or comma-separated ISO3 codes")
    ap.add_argument("--partner", help="'world', 'all', or comma-separated ISO3 codes")
    ap.add_argument("--out", help="Output directory")
    ap.add_argument("--delay", type=float, help="Seconds between API calls")
    ap.add_argument("--no-cache", action="store_true", help="Ignore the response cache")
    ap.add_argument("--list-codes", metavar="TERM",
                    help="Search HS product descriptions for TERM and exit")
    args = ap.parse_args()

    global PRODUCT_CODES
    if args.codes:
        PRODUCT_CODES = {c.strip(): "" for c in args.codes.split(",") if c.strip()}

    years_spec = parse_years_arg(args.years) if args.years else YEARS
    reporters_spec = parse_list_arg(args.reporters) if args.reporters else REPORTERS
    partner_spec = parse_list_arg(args.partner) if args.partner else PARTNER
    out_dir = pathlib.Path(args.out or OUTPUT_DIR)
    delay = args.delay if args.delay is not None else REQUEST_DELAY_SECONDS
    cache_dir = None if (args.no_cache or not USE_CACHE) else pathlib.Path(CACHE_DIR)

    api_key = os.environ.get(API_KEY_ENV_VAR) or None
    client = Client(api_key, delay, cache_dir)
    lookups = Lookups(client)

    if args.list_codes:
        term = args.list_codes.lower()
        hits = [(c, d) for c, d in sorted(lookups.commodities.items())
                if len(c) == 6 and term in d.lower()]
        if not hits:
            print(f"No HS 6-digit codes matched {args.list_codes!r}.")
        for code, desc in hits:
            print(f"{code}  {desc}")
        return 0

    if not PRODUCT_CODES:
        return print("PRODUCT_CODES is empty -- nothing to pull.") or 1

    years = resolve_years(years_spec)
    codes = sorted(PRODUCT_CODES)
    reporter_codes = lookups.reporter_codes(reporters_spec)
    partner_param = lookups.partner_param(partner_spec)
    # An explicit reporter list can exceed URL limits; "all" is sent as empty.
    reporters_arg = None if reporters_spec == "all" else reporter_codes

    print(f"\nSource      : UN Comtrade (the HS6 series WITS redistributes)")
    print(f"Auth        : {'API key from $' + API_KEY_ENV_VAR if api_key else 'keyless (500-row cap, throttled)'}")
    print(f"Products    : {', '.join(codes)}")
    print(f"Years       : {years[0]}-{years[-1]} ({len(years)} years)")
    print(f"Reporters   : {'all' if reporters_spec == 'all' else f'{len(reporter_codes)} selected'}")
    print(f"Partner     : {partner_spec if isinstance(partner_spec, str) else ','.join(partner_spec)}")
    print(f"Flows       : {', '.join(f'{k}={v}' for k, v in FLOWS.items())}")
    print(f"Output      : {out_dir.resolve()}\n")

    collected: dict[str, list[dict]] = {label: [] for label in FLOWS.values()}
    empty_years: list[int] = []

    for year in years:
        year_total = 0
        for flow_code, flow_label in FLOWS.items():
            raw = fetch_partition(client, year, flow_code, codes, reporters_arg, partner_param)
            rows = normalise(raw, lookups, flow_label)
            collected[flow_label].extend(rows)
            year_total += len(rows)
        if year_total:
            print(f"  {year}: {year_total:>6,} rows", flush=True)
        else:
            empty_years.append(year)

    if empty_years:
        print(f"  no data reported for: {', '.join(str(y) for y in empty_years)}")

    out_dir.mkdir(parents=True, exist_ok=True)
    for label in collected:
        collected[label].sort(key=sort_key)

    written = []
    for label, rows in collected.items():
        path = out_dir / f"lithium_{label.lower()}.csv"
        write_csv(path, rows)
        written.append((path, len(rows)))

    xlsx_path = out_dir / "lithium_trade.xlsx"
    xlsx_ok = write_xlsx(xlsx_path, collected)

    print("\nDone.")
    for path, count in written:
        print(f"  {path}  ({count:,} rows)")
    if xlsx_ok:
        print(f"  {xlsx_path}  (sheets: {', '.join(collected)})")
    else:
        print(f"  (skipped {xlsx_path.name} -- run 'pip install openpyxl' for the "
              f"two-sheet workbook; the CSVs above are complete)")
    print(f"\nAPI calls: {client.calls}   cache hits: {client.cache_hits}")
    return 0


if __name__ == "__main__":
    try:
        sys.exit(main())
    except KeyboardInterrupt:
        print("\nInterrupted. Cached responses are kept -- rerun to resume.")
        sys.exit(130)
